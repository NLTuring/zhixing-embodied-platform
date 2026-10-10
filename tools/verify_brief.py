# -*- coding: utf-8 -*-
"""校验生成的 DOCX：字体是否全为宋体、是否残留 Markdown 标记、结构是否完整。"""
import io
import os
import sys

from docx import Document
from docx.oxml.ns import qn

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
# 可选参数：argv[1] = 待校验 docx 路径，argv[2] = 报告输出路径
DOCX = (sys.argv[1] if len(sys.argv) > 1 else
        os.path.join(ROOT, '具身智能任务规划中间件与仿真评测平台-项目简介文档.docx'))
REPORT = (sys.argv[2] if len(sys.argv) > 2 else
          os.path.join(ROOT, 'tools', '_verify_report.txt'))
SONG = '宋体'
ATTRS = ('w:ascii', 'w:hAnsi', 'w:eastAsia', 'w:cs')

doc = Document(DOCX)

bad_font = []
stars = []
empty = []


def check_runs(paragraphs, where):
    for pi, para in enumerate(paragraphs):
        for run in para.runs:
            txt = run.text or ''
            if '**' in txt:
                stars.append((where, pi, txt[:60]))
            rPr = run._element.find(qn('w:rPr'))
            rFonts = rPr.find(qn('w:rFonts')) if rPr is not None else None
            if rFonts is None:
                if txt.strip():
                    bad_font.append((where, pi, 'NO rFonts', txt[:40]))
                continue
            for a in ATTRS:
                v = rFonts.get(qn(a))
                if v != SONG:
                    if txt.strip():
                        bad_font.append((where, pi, a, '%s | %s' % (v, txt[:40])))


check_runs(doc.paragraphs, 'body')

table_runs = 0
for ti, t in enumerate(doc.tables):
    for row in t.rows:
        for cell in row.cells:
            for para in cell.paragraphs:
                for run in para.runs:
                    table_runs += 1
                    if '**' in (run.text or ''):
                        stars.append(('table%d' % ti, 0, run.text[:60]))
check_runs([pp for t in doc.tables for row in t.rows
            for cell in row.cells for pp in cell.paragraphs], 'table')

for si, sec in enumerate(doc.sections):
    check_runs(sec.header.paragraphs, 'hdr%d' % si)
    check_runs(sec.footer.paragraphs, 'ftr%d' % si)

styles_bad = []
for st in doc.styles:
    try:
        rPr = st.element.find(qn('w:rPr'))
        if rPr is None:
            continue
        rFonts = rPr.find(qn('w:rFonts'))
        if rFonts is None:
            continue
        for a in ATTRS:
            v = rFonts.get(qn(a))
            if v is not None and v != SONG:
                styles_bad.append((st.name, a, v))
    except Exception:
        pass

headings = [(p.style.name, p.text) for p in doc.paragraphs
            if p.style.name.startswith('Heading')]

xml = doc.element.xml
has_toc = 'TOC \\o' in xml
has_updatefields = 'updateFields' in doc.settings.element.xml
sect_count = len(doc.sections)

lines = []
lines.append('=== DOCX 校验报告 ===')
lines.append('文件: %s' % os.path.basename(DOCX))
lines.append('大小: %.1f KB' % (os.path.getsize(DOCX) / 1024.0))
lines.append('段落数: %d   表格数: %d   表格内 run: %d   节数: %d'
             % (len(doc.paragraphs), len(doc.tables), table_runs, sect_count))
lines.append('目录域(TOC): %s   打开时自动更新域: %s' % (has_toc, has_updatefields))
lines.append('标题数: %d' % len(headings))
lines.append('')
lines.append('--- 非宋体 run: %d ---' % len(bad_font))
for b in bad_font[:40]:
    lines.append('  %s #%d [%s] %s' % (b[0], b[1], b[2], b[3]))
lines.append('--- 样式表非宋体: %d ---' % len(styles_bad))
for s in styles_bad[:20]:
    lines.append('  %s / %s = %s' % s)
lines.append('--- 残留 ** 标记: %d ---' % len(stars))
for s in stars[:20]:
    lines.append('  %s #%d %s' % s)
lines.append('')
lines.append('--- 标题结构 ---')
for st, tx in headings:
    lvl = st.replace('Heading ', '')
    lines.append('%s %s%s' % ('  ' * (int(lvl) - 1) if lvl.isdigit() else '', tx,
                              '' if lvl == '1' else ''))

report = '\n'.join(lines)
with io.open(REPORT, 'w', encoding='utf-8') as f:
    f.write(report)
sys.stdout.reconfigure(encoding='utf-8')
print(report)
